#!/usr/bin/env python3
"""
analysis_agent.py — Agent 4 of the notebook scheduling pipeline.

Parses siparisturu_v6.lst (the GAMS DISPLAY section) plus the job names and
qty values from siparisturu_v6.gms, then writes a 3-sheet Excel report:

  Sheet 1: "Job Schedule"   — all 198 jobs, sorted by tardiness (worst first)
  Sheet 2: "Machine Load"   — per-machine load and utilization
  Sheet 3: "Summary"        — KPI block

Exit codes:
    0  success
    1  siparisturu_v6.lst not found
    2  .lst exists but no DISPLAY section
    3  openpyxl not installed

Uses only Python standard library + openpyxl.
"""

from __future__ import annotations

import os
import re
import sys
from typing import Dict, List, Optional, Tuple

LST_FILE = "siparisturu_v6.lst"
GMS_FILE = "siparisturu_v6.gms"
OUTPUT_FILE = "schedule_output.xlsx"

ROUTE_DESCRIPTIONS = {
    1: "Wire Stitch (Bielom2 -> Shrink1 -> MSK)",
    2: "Plastic Spiral via Plasticol",
    3: "Plastic Spiral via Kore",
    4: "Double Metal via Autobind",
    5: "Double Metal via CWH",
    6: "Manual Spiral",
    7: "Thread Stitch",
}

MACHINE_CAPACITIES = {
    "Bielom1":    4800,
    "Bielom2":    4800,
    "Shrink1":    2400,
    "Shrink2":    2400,
    "MSK":        2400,
    "PolarBicak": 1200,
    "Kugler":     1200,
    "Plasticol":  1200,
    "Kore":       1200,
    "Autobind":   1200,
    "CWH":        1200,
    "ManSpiral":  1200,
    "Juki":       1200,
}


# ---------------------------------------------------------------------------
# .gms helpers (reused from data_agent's logic; kept self-contained per spec)
# ---------------------------------------------------------------------------

def parse_job_names(gms_text: str) -> Dict[str, str]:
    m = re.search(r"Set\s+i\s+isler\s*/(.*?)/", gms_text, re.DOTALL | re.IGNORECASE)
    if not m:
        return {}
    block = m.group(1)
    names: Dict[str, str] = {}
    for line in block.splitlines():
        line = line.strip()
        if not line or line.startswith("*"):
            continue
        match = re.match(r'(I\d+)\s+"([^"]*)"', line)
        if match:
            names[match.group(1)] = match.group(2)
    return names


def parse_qty(gms_text: str) -> Dict[str, int]:
    m = re.search(r"Parameter\s+qty\(i\)[^/]*/(.*?)/", gms_text, re.DOTALL | re.IGNORECASE)
    if not m:
        return {}
    block = m.group(1)
    qty: Dict[str, int] = {}
    for line in block.splitlines():
        line = line.strip()
        if not line or line.startswith("*"):
            continue
        match = re.match(r"(I\d+)\s+([\d.]+)", line)
        if match:
            qty[match.group(1)] = int(float(match.group(2)))
    return qty


def extract_customer_binding_ruling(job_name: str) -> Tuple[str, str, str]:
    parts = job_name.split("-")
    if len(parts) >= 5:
        return parts[0].strip(), parts[1].strip(), parts[-1].strip()
    if len(parts) >= 2:
        return parts[0].strip(), parts[1].strip(), parts[-1].strip()
    return job_name, "", ""


# ---------------------------------------------------------------------------
# .lst parser
# ---------------------------------------------------------------------------

# Matches a parameter (or variable.l) table header line.
# Examples:  "----    488 PARAMETER rota_atama  route assigned to each job"
#            "----    488 VARIABLE  X.L         level values"
PARAM_HEADER = re.compile(
    r"^----\s*\d+\s+(?:PARAMETER|VARIABLE)\s+(\S+?)\s*($|\s+[^=].*$)"
)
# Matches a scalar PARAMETER/VARIABLE line:
#   "----  488 PARAMETER Z_obj.l   =  123.4"
#   "----  488 VARIABLE  Cmax.L    =  987.6"
PARAM_SCALAR = re.compile(
    r"^----\s*\d+\s+(?:PARAMETER|VARIABLE)\s+(\S+?)\s*=\s*([\-\d.eE+]+)"
)
# Other section delimiters in the .lst file:
SECTION_BREAK = re.compile(r"^----")
# An element/value pair in a parameter table:  "I01  2.000"
ELEMENT_VALUE = re.compile(r"([A-Za-z_]\w*?)\s+([\-\d.eE+]+)")


def _normalize_name(raw: str) -> str:
    """Lowercase a GAMS symbol name and strip a trailing .l/.L attribute."""
    n = raw.lower().rstrip(".")
    if n.endswith(".l"):
        n = n[:-2]
    return n


def parse_lst(lst_path: str) -> Tuple[Dict[str, Dict[str, float]], Dict[str, float], bool]:
    """Parse the GAMS .lst file.

    Returns: (tables, scalars, has_display)
      tables  -- dict {param_name_lower: {element_name: float}}
      scalars -- dict {param_name_lower: float}
      has_display -- True if at least one PARAMETER line was seen
    """
    tables: Dict[str, Dict[str, float]] = {}
    scalars: Dict[str, float] = {}
    has_display = False

    current_name: Optional[str] = None
    current_dict: Optional[Dict[str, float]] = None

    def finalize() -> None:
        nonlocal current_name, current_dict
        if current_name is not None and current_dict is not None:
            # Merge: later occurrences just update earlier values
            tables.setdefault(current_name, {}).update(current_dict)
        current_name = None
        current_dict = None

    with open(lst_path, "r", encoding="utf-8", errors="replace") as f:
        for raw in f:
            line = raw.rstrip("\n")

            # Scalar PARAMETER line takes precedence (matches "name = value")
            ms = PARAM_SCALAR.match(line)
            if ms:
                finalize()
                # Normalize: lowercase + strip ".l" suffix (GAMS writes Z_obj.L / Z_obj.l)
                key = _normalize_name(ms.group(1))
                try:
                    scalars[key] = float(ms.group(2))
                    has_display = True
                except ValueError:
                    pass
                continue

            # Table header line:  "----   488 PARAMETER rota_atama [description]"
            mh = PARAM_HEADER.match(line)
            if mh:
                finalize()
                current_name = _normalize_name(mh.group(1))
                current_dict = {}
                has_display = True
                continue

            # A non-PARAMETER section break ends the current table.
            if SECTION_BREAK.match(line) and current_name is not None:
                finalize()
                continue

            # Skip page-break / banner / empty lines while inside a table.
            if current_name is None:
                continue

            # Inside a table body: extract (element, value) pairs.
            stripped = line.strip()
            if not stripped:
                # blank lines inside a table do not end it (continuation)
                continue
            if stripped.startswith("GAMS ") or stripped.startswith("Compilation"):
                # GAMS page header inside a long table; ignore
                continue
            if stripped.startswith("+"):
                # Continuation marker in GAMS output -- skip the leading symbol
                stripped = stripped.lstrip("+").strip()

            for elem, val in ELEMENT_VALUE.findall(stripped):
                try:
                    current_dict[elem] = float(val)
                except ValueError:
                    continue

        finalize()

    return tables, scalars, has_display


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------

def build_job_table(
    job_names: Dict[str, str],
    qty: Dict[str, int],
    tables: Dict[str, Dict[str, float]],
    scalars: Dict[str, float],
) -> List[Dict[str, object]]:
    rota = tables.get("rota_atama", {})
    tamamlanma = tables.get("tamamlanma", {})
    teslim_gun = tables.get("teslim_gun", {})
    gecikme = tables.get("gecikme", {})
    gecikme_gun = tables.get("gecikme_gun", {})
    d_min = tables.get("d", {})

    rows: List[Dict[str, object]] = []
    job_ids = sorted(job_names.keys(), key=lambda s: int(s[1:]))
    for j in job_ids:
        name = job_names.get(j, "")
        customer, binding, ruling = extract_customer_binding_ruling(name)
        route_num = int(round(rota.get(j, 0)))
        completion_min = tamamlanma.get(j, 0.0)
        completion_days = teslim_gun.get(j, completion_min / 1200.0 if completion_min else 0.0)
        tardiness_min = gecikme.get(j, 0.0)
        tardiness_days = gecikme_gun.get(j, tardiness_min / 1200.0 if tardiness_min else 0.0)
        due_minutes = d_min.get(j, 0.0)
        due_days = due_minutes / 1200.0 if due_minutes else 0.0

        if completion_min == 0:
            status = "UNSCHEDULED"
        elif tardiness_min > 0.01:
            status = "LATE"
        else:
            status = "ON TIME"

        rows.append({
            "job_id": j,
            "job_name": name,
            "customer": customer,
            "binding_type": binding,
            "ruling": ruling,
            "route_num": route_num,
            "route_desc": ROUTE_DESCRIPTIONS.get(route_num, "(unassigned)"),
            "qty": qty.get(j, 0),
            "due_minutes": due_minutes,
            "due_days": due_days,
            "completion_min": completion_min,
            "completion_days": completion_days,
            "tardiness_min": tardiness_min,
            "tardiness_days": tardiness_days,
            "status": status,
        })

    # Sort: tardiness desc, then due_days asc for ties
    rows.sort(key=lambda r: (-float(r["tardiness_days"]), float(r["due_days"])))
    return rows


def write_excel(
    output_path: str,
    rows: List[Dict[str, object]],
    tables: Dict[str, Dict[str, float]],
    scalars: Dict[str, float],
) -> Dict[str, object]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ----- Sheet 1: Job Schedule -----
    ws = wb.active
    ws.title = "Job Schedule"

    headers = [
        "job_id", "job_name", "customer", "binding_type", "ruling",
        "route_num", "route_desc", "qty",
        "due_minutes", "due_days",
        "completion_min", "completion_days",
        "tardiness_min", "tardiness_days", "status",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    gray_fill = PatternFill("solid", fgColor="D9D9D9")
    bold = Font(bold=True)

    for ridx, row in enumerate(rows, start=2):
        for cidx, key in enumerate(headers, start=1):
            ws.cell(row=ridx, column=cidx, value=row[key])
        status_cell = ws.cell(row=ridx, column=headers.index("status") + 1)
        if row["status"] == "LATE":
            status_cell.fill = red_fill
        elif row["status"] == "ON TIME":
            status_cell.fill = green_fill
        else:
            status_cell.fill = gray_fill

        if float(row["tardiness_days"]) > 1:
            for cidx in range(1, len(headers) + 1):
                ws.cell(row=ridx, column=cidx).font = bold

    # number formatting for numeric columns
    int_cols = {"qty", "route_num"}
    one_decimal_cols = {"due_minutes", "completion_min", "tardiness_min"}
    two_decimal_cols = {"due_days", "completion_days", "tardiness_days"}
    for cidx, key in enumerate(headers, start=1):
        if key in int_cols:
            fmt = "0"
        elif key in one_decimal_cols:
            fmt = "0.0"
        elif key in two_decimal_cols:
            fmt = "0.00"
        else:
            fmt = None
        if fmt:
            for ridx in range(2, len(rows) + 2):
                ws.cell(row=ridx, column=cidx).number_format = fmt

    # auto-ish column widths
    widths = {
        "job_id": 8, "job_name": 48, "customer": 14, "binding_type": 22,
        "ruling": 10, "route_num": 9, "route_desc": 36, "qty": 12,
        "due_minutes": 12, "due_days": 10,
        "completion_min": 14, "completion_days": 14,
        "tardiness_min": 14, "tardiness_days": 14, "status": 12,
    }
    for cidx, key in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(cidx)].width = widths.get(key, 14)
    ws.freeze_panes = "A2"

    # ----- Sheet 2: Machine Load -----
    ws2 = wb.create_sheet("Machine Load")
    yuk = tables.get("makine_yuk", {})
    yuk_gun = tables.get("makine_yuk_gun", {})

    cmax = scalars.get("cmax") or 0.0
    total_days = (cmax / 1200.0) if cmax else 0.0

    machine_headers = [
        "Machine", "Load (min)", "Load (days)",
        "Capacity (min/day)", "Utilization %",
    ]
    for cidx, h in enumerate(machine_headers, start=1):
        cell = ws2.cell(row=1, column=cidx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    machine_rows: List[Tuple[str, float, float, int, float]] = []
    for m, cap in MACHINE_CAPACITIES.items():
        load_min = float(yuk.get(m, 0.0))
        load_days = float(yuk_gun.get(m, load_min / 1200.0 if load_min else 0.0))
        denom = cap * total_days if total_days > 0 else 0.0
        util_pct = (load_min / denom * 100.0) if denom > 0 else 0.0
        machine_rows.append((m, load_min, load_days, cap, util_pct))

    machine_rows.sort(key=lambda x: -x[4])  # by utilization desc
    for ridx, (m, load_min, load_days, cap, util_pct) in enumerate(machine_rows, start=2):
        ws2.cell(row=ridx, column=1, value=m)
        ws2.cell(row=ridx, column=2, value=round(load_min, 1)).number_format = "0.0"
        ws2.cell(row=ridx, column=3, value=round(load_days, 3)).number_format = "0.000"
        ws2.cell(row=ridx, column=4, value=cap).number_format = "0"
        ws2.cell(row=ridx, column=5, value=round(util_pct, 2)).number_format = "0.00"

    for cidx, w in enumerate([16, 14, 12, 16, 16], start=1):
        ws2.column_dimensions[get_column_letter(cidx)].width = w
    ws2.freeze_panes = "A2"

    # ----- Sheet 3: Summary KPI -----
    ws3 = wb.create_sheet("Summary")
    total_jobs = len(rows)
    on_time_count = sum(1 for r in rows if r["status"] == "ON TIME")
    unscheduled_count = sum(1 for r in rows if r["status"] == "UNSCHEDULED")
    late_count_calc = sum(1 for r in rows if r["status"] == "LATE")
    geciken_is = int(round(scalars.get("geciken_is_adedi", late_count_calc)))
    on_time_rate = (on_time_count / total_jobs * 100.0) if total_jobs else 0.0

    most_loaded = machine_rows[0] if machine_rows else ("n/a", 0, 0, 0, 0)
    worst_late = rows[0] if rows else None
    z_obj = scalars.get("z_obj")

    summary_block = [
        ("Total jobs scheduled", total_jobs),
        ("Jobs on time",         on_time_count),
        ("Jobs late",            geciken_is),
        ("Jobs unscheduled",     unscheduled_count),
        ("On-time rate",         f"{on_time_rate:.1f}%"),
        ("Objective value (Z)",  z_obj if z_obj is not None else "n/a"),
        (
            "Makespan",
            f"{cmax:.1f} min  /  {total_days:.2f} days" if cmax else "n/a",
        ),
        (
            "Most loaded machine",
            f"{most_loaded[0]} at {most_loaded[4]:.1f}% (load {most_loaded[1]:.0f} min)",
        ),
        (
            "Most late job",
            (
                f"{worst_late['job_id']} -- {worst_late['job_name']}, "
                f"{float(worst_late['tardiness_days']):.2f} days late"
                if worst_late and float(worst_late["tardiness_days"]) > 0
                else "none"
            ),
        ),
    ]
    ws3.cell(row=1, column=1, value="Metric").font = header_font
    ws3.cell(row=1, column=1).fill = header_fill
    ws3.cell(row=1, column=2, value="Value").font = header_font
    ws3.cell(row=1, column=2).fill = header_fill
    for ridx, (label, value) in enumerate(summary_block, start=2):
        ws3.cell(row=ridx, column=1, value=label).font = bold
        ws3.cell(row=ridx, column=2, value=value)
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 64

    wb.save(output_path)

    return {
        "on_time_count": on_time_count,
        "late_count": geciken_is,
        "unscheduled": unscheduled_count,
        "worst_late": worst_late,
        "cmax_days": total_days,
    }


def main() -> int:
    if not os.path.exists(LST_FILE):
        print(f"ERROR: {LST_FILE} not found in working directory")
        return 1
    if not os.path.exists(GMS_FILE):
        print(f"ERROR: {GMS_FILE} not found in working directory")
        return 1

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("ERROR: openpyxl not installed -- run: pip install openpyxl")
        return 3

    with open(GMS_FILE, "r", encoding="utf-8") as f:
        gms_text = f.read()
    job_names = parse_job_names(gms_text)
    qty = parse_qty(gms_text)

    tables, scalars, has_display = parse_lst(LST_FILE)
    if not has_display:
        print(f"ERROR: {LST_FILE} contains no DISPLAY section -- solve may have failed")
        return 2

    rows = build_job_table(job_names, qty, tables, scalars)
    result = write_excel(OUTPUT_FILE, rows, tables, scalars)

    cmax_days = result["cmax_days"]
    worst = result["worst_late"]
    worst_str = (
        f"{worst['job_id']}, {float(worst['tardiness_days']):.2f} days"
        if worst and float(worst["tardiness_days"]) > 0
        else "none"
    )

    print()
    print("ANALYSIS AGENT COMPLETE")
    print(f"Output: {OUTPUT_FILE}")
    print(f"Jobs on time: {result['on_time_count']} / {len(rows)}")
    print(f"Most late job: {worst_str}")
    print(f"Makespan: {cmax_days:.2f} working days")
    return 0


if __name__ == "__main__":
    sys.exit(main())
